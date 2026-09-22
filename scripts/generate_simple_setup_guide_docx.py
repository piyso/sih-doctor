#!/usr/bin/env python3
"""
AIIA Hospital OS - Simple Step-by-Step Setup & 3-Minute Live Demo Guide Generator
Generates: AIIA_Hospital_OS_Simple_Setup_and_Demo_Guide.docx
Clear, plain English, zero jargon, institutional formatting with tables and visual boxes.
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR_26047 = "/Users/piyushkumar/Desktop/SIH/26047"
OUTPUT_DIR_ROOT = "/Users/piyushkumar/Desktop/SIH"
FILENAME = "AIIA_Hospital_OS_Simple_Setup_and_Demo_Guide.docx"
DOCX_PATH_26047 = os.path.join(OUTPUT_DIR_26047, FILENAME)
DOCX_PATH_ROOT = os.path.join(OUTPUT_DIR_ROOT, FILENAME)

os.makedirs(OUTPUT_DIR_26047, exist_ok=True)

# Colors
COLOR_PRIMARY = RGBColor(6, 78, 59)       # Deep Forest Green (#064E3B)
COLOR_SECONDARY = RGBColor(2, 132, 199)   # Sky Blue (#0284C7)
COLOR_TEXT_DARK = RGBColor(15, 23, 42)    # Slate 900
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)# Slate 500
COLOR_ACCENT = RGBColor(180, 83, 9)       # Warm Amber (#B45309)
COLOR_DANGER = RGBColor(220, 38, 38)      # Alert Red

HEX_PRIMARY = "064E3B"
HEX_SECONDARY = "0284C7"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_SUCCESS_BG = "F0FDF4"
HEX_SUCCESS_BORDER = "BBF7D0"
HEX_ALERT_BG = "FEF2F2"
HEX_ALERT_BORDER = "FECACA"
HEX_ACCENT_BG = "FFFBEB"
HEX_ACCENT_BORDER = "FDE68A"

def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(shd)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    
    borders = {'w:top': top, 'w:bottom': bottom, 'w:left': left, 'w:right': right}
    for b_name, b_val in borders.items():
        if b_val is not None:
            el = parse_xml(f'<{b_name} {nsdecls("w")} w:val="{b_val.get("val", "single")}" w:sz="{b_val.get("sz", 4)}" w:space="0" w:color="{b_val.get("color", "auto")}"/>')
            tcBorders.append(el)

def add_callout(doc, text_list, title="NOTICE", bg_hex=HEX_SUCCESS_BG, border_color="10B981"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_borders(cell, 
                     left={'val': 'single', 'sz': 24, 'color': border_color},
                     top={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                     bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                     right={'val': 'single', 'sz': 4, 'color': HEX_BORDER})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY if border_color != "DC2626" else COLOR_DANGER
    
    for item in text_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        run_item = p_item.add_run(item)
        run_item.font.name = "Arial"
        run_item.font.size = Pt(9.5)
        run_item.font.color.rgb = COLOR_TEXT_DARK
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_document():
    print(f"Creating simplified setup guide: {DOCX_PATH_26047}")
    doc = docx.Document()
    
    # Set standard margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header setup
        header = section.header
        header.is_linked_to_previous = False
        tbl_h = header.add_table(rows=1, cols=2, width=Inches(6.9))
        tbl_h.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_h_left = tbl_h.cell(0, 0)
        cell_h_right = tbl_h.cell(0, 1)
        
        p_hl = cell_h_left.paragraphs[0]
        p_hl.paragraph_format.space_after = Pt(0)
        r_hl = p_hl.add_run("ALL INDIA INSTITUTE OF AYURVEDA · PS ID 26047")
        r_hl.font.name = "Arial"
        r_hl.font.size = Pt(8.5)
        r_hl.font.color.rgb = COLOR_PRIMARY
        r_hl.font.bold = True
        
        p_hr = cell_h_right.paragraphs[0]
        p_hr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hr.paragraph_format.space_after = Pt(0)
        r_hr = p_hr.add_run("SIMPLE SETUP & DEMO GUIDE")
        r_hr.font.name = "Arial"
        r_hr.font.size = Pt(8.5)
        r_hr.font.color.rgb = COLOR_TEXT_MUTED
        
        # Footer setup
        footer = section.footer
        footer.is_linked_to_previous = False
        tbl_f = footer.add_table(rows=1, cols=2, width=Inches(6.9))
        tbl_f.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_fl = tbl_f.cell(0, 0)
        cell_fr = tbl_f.cell(0, 1)
        
        p_fl = cell_fl.paragraphs[0]
        p_fl.paragraph_format.space_after = Pt(0)
        r_fl = p_fl.add_run("Sovereign Air-Gapped Hospital OS · Smart India Hackathon 2026")
        r_fl.font.name = "Arial"
        r_fl.font.size = Pt(8.5)
        r_fl.font.color.rgb = COLOR_TEXT_MUTED
        
        p_fr = cell_fr.paragraphs[0]
        p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_fr.paragraph_format.space_after = Pt(0)
        r_fr = p_fr.add_run("Confidential · Team Agastya Sutra")
        r_fr.font.name = "Arial"
        r_fr.font.size = Pt(8.5)
        r_fr.font.color.rgb = COLOR_TEXT_MUTED

    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("🏛️ AIIA Sovereign Hospital OS\nSimple Step-by-Step Setup & 3-Minute Demo Playbook")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Problem Statement 26047 · Ministry of Ayush & MoHFW · Smart India Hackathon 2026")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Callout: The Golden Rule (No internet needed!)
    add_callout(
        doc,
        [
            "• Zero Internet Required: The entire hospital system runs 100% locally on your PC in complete Airplane Mode.",
            "• Zero Database Configuration: Powered by embedded SQLite WAL. No Docker or cloud servers to set up.",
            "• One Address to Remember: On your computer, always open: http://localhost:5173/"
        ],
        title="THE GOLDEN RULE — 100% AIR-GAPPED SOVEREIGN OPERATION",
        bg_hex=HEX_SUCCESS_BG,
        border_color="059669"
    )

    # -------------------------------------------------------------
    # SECTION 1: HOW TO START THE SYSTEM (2 STEPS)
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("Part 1: How to Start the System on Any Laptop (Only 2 Steps)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    # Step 1 Box
    tbl_s1 = doc.add_table(rows=1, cols=2)
    tbl_s1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s1.autofit = False
    
    c_num = tbl_s1.cell(0, 0)
    c_num.width = Inches(1.1)
    set_cell_shading(c_num, HEX_PRIMARY)
    set_cell_margins(c_num, 120, 120, 120, 120)
    p_num = c_num.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_num = p_num.add_run("STEP\n1")
    r_num.font.name = "Arial"
    r_num.font.size = Pt(18)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(255, 255, 255)
    
    c_desc = tbl_s1.cell(0, 1)
    c_desc.width = Inches(5.7)
    set_cell_shading(c_desc, HEX_LIGHT_BG)
    set_cell_margins(c_desc, 120, 120, 140, 140)
    set_cell_borders(c_desc, top={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                             bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                             right={'val': 'single', 'sz': 4, 'color': HEX_BORDER})
    
    p_d = c_desc.paragraphs[0]
    p_d.paragraph_format.space_after = Pt(3)
    r_d1 = p_d.add_run("Start the System Launcher:\n")
    r_d1.font.name = "Arial"
    r_d1.font.size = Pt(11)
    r_d1.font.bold = True
    r_d1.font.color.rgb = COLOR_TEXT_DARK
    
    p_w = c_desc.add_paragraph()
    p_w.paragraph_format.space_after = Pt(2)
    r_w = p_w.add_run("• On Windows: Open the folder and double-click start.bat (or run.bat).\n• On Mac / Linux: Open Terminal in the folder and run ./start.sh (or npm start).")
    r_w.font.name = "Arial"
    r_w.font.size = Pt(9.5)
    r_w.font.color.rgb = COLOR_TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Step 2 Box
    tbl_s2 = doc.add_table(rows=1, cols=2)
    tbl_s2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s2.autofit = False
    
    c_num2 = tbl_s2.cell(0, 0)
    c_num2.width = Inches(1.1)
    set_cell_shading(c_num2, HEX_SECONDARY)
    set_cell_margins(c_num2, 120, 120, 120, 120)
    p_num2 = c_num2.paragraphs[0]
    p_num2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_num2 = p_num2.add_run("STEP\n2")
    r_num2.font.name = "Arial"
    r_num2.font.size = Pt(18)
    r_num2.font.bold = True
    r_num2.font.color.rgb = RGBColor(255, 255, 255)
    
    c_desc2 = tbl_s2.cell(0, 1)
    c_desc2.width = Inches(5.7)
    set_cell_shading(c_desc2, HEX_LIGHT_BG)
    set_cell_margins(c_desc2, 120, 120, 140, 140)
    set_cell_borders(c_desc2, top={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                              bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                              right={'val': 'single', 'sz': 4, 'color': HEX_BORDER})
    
    p_d2 = c_desc2.paragraphs[0]
    p_d2.paragraph_format.space_after = Pt(3)
    r_d2 = p_d2.add_run("Open Your Web Browser:\n")
    r_d2.font.name = "Arial"
    r_d2.font.size = Pt(11)
    r_d2.font.bold = True
    r_d2.font.color.rgb = COLOR_TEXT_DARK
    
    p_w2 = c_desc2.add_paragraph()
    p_w2.paragraph_format.space_after = Pt(2)
    r_w2 = p_w2.add_run("Open Google Chrome, Edge, Safari, or Brave and go to:\n")
    r_w2.font.name = "Arial"
    r_w2.font.size = Pt(9.5)
    r_w2.font.color.rgb = COLOR_TEXT_DARK
    
    r_link = p_w2.add_run("http://localhost:5173/\n")
    r_link.font.name = "Arial"
    r_link.font.size = Pt(12)
    r_link.font.bold = True
    r_link.font.color.rgb = COLOR_SECONDARY
    
    r_w2_sub = p_w2.add_run("The AIIA Hospital OS Master Gateway will appear immediately!")
    r_w2_sub.font.name = "Arial"
    r_w2_sub.font.size = Pt(9.5)
    r_w2_sub.font.color.rgb = COLOR_TEXT_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Step 3 (Optional Mobile)
    add_callout(
        doc,
        [
            "1. Connect your phone or tablet to the same Wi-Fi or mobile hotspot as your laptop.",
            "2. Look at your computer's terminal screen. It will show a line with your exact numbers, for example:",
            "   Mobile Wi-Fi Access: http://10.241.19.238:5173",
            "3. Open the browser on your phone and type that exact address shown on your screen.",
            "4. Both your phone and laptop will now be synchronized in real-time!"
        ],
        title="📱 OPTIONAL STEP 3: HOW TO OPEN ON YOUR MOBILE PHONE",
        bg_hex=HEX_ACCENT_BG,
        border_color="D97706"
    )

    # -------------------------------------------------------------
    # SECTION 2: THE 6 HOSPITAL SCREENS
    # -------------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("Part 2: The 6 Hospital Workstation Screens (Explained in Simple Words)")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    tbl_screens = doc.add_table(rows=7, cols=3)
    tbl_screens.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_screens.autofit = False
    
    headers = ["Key", "Workstation Screen", "What It Does in Simple Words"]
    col_widths = [Inches(0.8), Inches(2.2), Inches(3.8)]
    
    for c_idx, h_text in enumerate(headers):
        cell = tbl_screens.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_shading(cell, HEX_PRIMARY)
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    screens_data = [
        ("1", "Patient MediKiosk\n(Lobby 32\" Touch)", "Patients enter their name, touch where it hurts on a 3D body map, answer 3 quick questions, and receive an instant token slip with an ABHA QR code."),
        ("2", "Doctor Clinical Cockpit\n(Chamber Desktop)", "Doctor sees the patient file pre-populated, turns on the microphone for live voice-to-text scribing, catches lethal drug interactions, and prints the Rx slip."),
        ("3", "Pharmacy POS\n(Dispensary Counter)", "Pharmacist scans the prescription barcode, verifies medicine safety, and prints peel-and-stick labels with dosage times."),
        ("4", "Frontline ASHA Field App\n(Rural Tablet)", "Simple tablet app for village healthcare workers visiting remote areas. Works 100% offline without any internet connection."),
        ("5", "Command Center NOC\n(Director's Wall)", "Hospital Director's live screen showing real-time patient queues, doctor consultation pacing, and disease outbreak alerts."),
        ("6", "System Defense Matrix\n(Audit Screen)", "Shows judges the mathematical privacy proofs (zk-SNARKs) and verifies that ZERO patient data leaves the hospital.")
    ]
    
    for r_idx, (k, title, desc) in enumerate(screens_data, start=1):
        bg = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
        
        c0 = tbl_screens.cell(r_idx, 0)
        c0.width = col_widths[0]
        set_cell_shading(c0, bg)
        set_cell_margins(c0, 80, 80, 80, 80)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(k)
        r0.font.name = "Arial"
        r0.font.size = Pt(11)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_SECONDARY
        
        c1 = tbl_screens.cell(r_idx, 1)
        c1.width = col_widths[1]
        set_cell_shading(c1, bg)
        set_cell_margins(c1, 80, 80, 80, 80)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(title)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_DARK
        
        c2 = tbl_screens.cell(r_idx, 2)
        c2.width = col_widths[2]
        set_cell_shading(c2, bg)
        set_cell_margins(c2, 80, 80, 80, 80)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(desc)
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        r2.font.color.rgb = COLOR_TEXT_DARK
        
        for c in (c0, c1, c2):
            set_cell_borders(c, top={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                                bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                                left={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                                right={'val': 'single', 'sz': 4, 'color': HEX_BORDER})

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 3: THE 3-MINUTE LIVE DEMO
    # -------------------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("Part 3: The 3-Minute Live Presentation Script (How to Wow Judges)")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_PRIMARY

    demo_steps = [
        ("Step 1: Patient Self-Intake at Kiosk (Press 1)",
         "• Open Patient MediKiosk (http://localhost:5173/?mode=kiosk).\n"
         "• Type Patient Name: Rajesh Sharma, Age: 45, Gender: Male.\n"
         "• Tap on the Chest on the 3D body mannequin and select 'Severe chest pain radiating to left arm'.\n"
         "• Click 'Generate OPD Token'. Show the judges the instant thermal slip with the official Indian National Emblem and ABHA QR Code!"),
        
        ("Step 2: Doctor Room & Live Ambient Microphone (Press 2)",
         "• Switch to Doctor Desk (http://localhost:5173/?mode=doctor).\n"
         "• Show that Rajesh Sharma is already in the live queue with an Emergency Red Flag alert!\n"
         "• Click on his name. His full pre-intake case sheet is already filled out with zero doctor typing.\n"
         "• Click 'Start Ambient Recording' and speak into your laptop microphone to show live voice transcription."),
        
        ("Step 3: Trigger the Lethal Herb-Drug Safety Alarm",
         "• In the Prescriber, add an Allopathic medicine: Warfarin (Blood thinner).\n"
         "• Then add an Ayurvedic medicine: Yogaraja Guggulu.\n"
         "• WATCH THE SCREEN: A bright red safety alert pops up instantly warning of severe internal bleeding risk!\n"
         "• Explain: 'Our system catches dangerous drug-herb clashes in under 1 millisecond.'\n"
         "• Replace Yogaraja Guggulu with safe medicine: Shallaki."),
        
        ("Step 4: Statutory Prescription Slip (Press Spacebar)",
         "• Press Spacebar or click 'Complete & Print Rx'.\n"
         "• A complete, statutory AIIA Government Prescription Slip is generated with official Ayush dosage directions, doctor credentials, and cryptographic verification!"),
        
        ("Step 5: Pharmacy & Command Center (Press 3 and 5)",
         "• Switch to Pharmacy POS (3) to show barcode scanning.\n"
         "• Switch to Command Center NOC (5) to show the live hospital director dashboard.")
    ]

    for title, body in demo_steps:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(6)
        p_st.paragraph_format.space_after = Pt(2)
        r_st = p_st.add_run(f"▶ {title}")
        r_st.font.name = "Arial"
        r_st.font.size = Pt(11)
        r_st.font.bold = True
        r_st.font.color.rgb = COLOR_PRIMARY
        
        p_sb = doc.add_paragraph()
        p_sb.paragraph_format.space_after = Pt(4)
        r_sb = p_sb.add_run(body)
        r_sb.font.name = "Arial"
        r_sb.font.size = Pt(9.5)
        r_sb.font.color.rgb = COLOR_TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 4: TROUBLESHOOTING & FAQ
    # -------------------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("Part 4: Frequently Asked Questions & Quick Help")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY

    faqs = [
        ("Q: How do I stop the servers when I am finished?",
         "Go to the black terminal window and press 'Ctrl + C' on your keyboard."),
        ("Q: What if the terminal says 'port 8001 or 5173 is in use'?",
         "Our startup script (start.sh or start.bat) automatically cleans and frees these ports. Simply close any extra terminal windows and double-click start.bat (or run ./start.sh) again."),
        ("Q: Does this require internet during the presentation?",
         "No. The entire system is 100% sovereign and air-gapped. You can demonstrate the entire hospital system with your Wi-Fi completely turned off (Airplane Mode)!")
    ]

    for q, a in faqs:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(1)
        r_q = p_q.add_run(q)
        r_q.font.name = "Arial"
        r_q.font.size = Pt(10)
        r_q.font.bold = True
        r_q.font.color.rgb = COLOR_SECONDARY
        
        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_after = Pt(4)
        r_a = p_a.add_run(a)
        r_a.font.name = "Arial"
        r_a.font.size = Pt(9.5)
        r_a.font.color.rgb = COLOR_TEXT_DARK

    # Save document in both locations
    doc.save(DOCX_PATH_26047)
    print(f"✓ Saved to: {DOCX_PATH_26047}")
    shutil.copy2(DOCX_PATH_26047, DOCX_PATH_ROOT)
    print(f"✓ Copied to: {DOCX_PATH_ROOT}")

if __name__ == "__main__":
    generate_document()
